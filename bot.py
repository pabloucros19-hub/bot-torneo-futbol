import os
import re
import logging
import webbrowser
from datetime import datetime
import openpyxl
from telegram import Update
from telegram.ext import ApplicationBuilder, ContextTypes, MessageHandler, filters

# Librerías para la API oficial de Google Drive
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload

# Configurar logging
logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)

EXCEL_FILE = "torneo_finanzas.xlsx"
SCOPES = ['https://www.googleapis.com/auth/drive.file']

# ----------------------------------------------------
# RESPALDO AUTOMÁTICO EN GOOGLE DRIVE
# ----------------------------------------------------
def respaldar_en_drive():
    try:
        chrome_path = "C:/Program Files/Google/Chrome/Application/chrome.exe"
        if os.path.exists(chrome_path):
            webbrowser.register('chrome', None, webbrowser.BackgroundBrowser(chrome_path))
            browser_to_use = 'chrome'
        else:
            chrome_path_alt = os.path.expanduser("~") + "/AppData/Local/Google/Chrome/Application/chrome.exe"
            if os.path.exists(chrome_path_alt):
                webbrowser.register('chrome', None, webbrowser.BackgroundBrowser(chrome_path_alt))
                browser_to_use = 'chrome'
            else:
                browser_to_use = None

        creds = None
        if os.path.exists('token.json'):
            creds = Credentials.from_authorized_user_file('token.json', SCOPES)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                flow = InstalledAppFlow.from_client_secrets_file('credentials.json', SCOPES)
                if browser_to_use:
                    webbrowser.get(browser_to_use)
                creds = flow.run_local_server(port=0)
            with open('token.json', 'w') as token:
                token.write(creds.to_json())

        service = build('drive', 'v3', credentials=creds)

        page_token = None
        file_id = None
        while True:
            response = service.files().list(
                q=f"name='{EXCEL_FILE}' and trashed=false",
                spaces='drive',
                fields='files(id, name)',
                pageToken=page_token
            ).execute()
            for file in response.get('files', []):
                file_id = file.get('id')
            page_token = response.get('nextPageToken', None)
            if page_token is None:
                break

        media = MediaFileUpload(EXCEL_FILE, resumable=True)
        if file_id:
            service.files().update(fileId=file_id, media_body=media).execute()
            logging.info("☁️ Respaldo en Google Drive actualizado con éxito.")
        else:
            file_metadata = {'name': EXCEL_FILE}
            service.files().create(body=file_metadata, media_body=media, fields='id').execute()
            logging.info("☁️ Archivo de respaldo creado en Google Drive por primera vez.")
    except Exception as e:
        logging.error(f"⚠️ Error al respaldar en Google Drive: {e}")

# ----------------------------------------------------
# INICIALIZACIÓN DE EXCEL
# ----------------------------------------------------
def inicializar_excel():
    columnas_pagos = [
        "Fecha", "Equipo Local", "Pago Local", 
        "Equipo Visitante", "Pago Visitante", 
        "Efectivo Total", "Nequi Total", 
        "Ingreso Total", "Descuento Fijo", "Caja Neta"
    ]
    
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws_pagos = wb.active
        ws_pagos.title = "Pagos Torneo"
        ws_pagos.append(columnas_pagos)
        
        ws_tarjetas = wb.create_sheet(title="Tarjetas")
        ws_tarjetas.append(["Fecha", "Equipo", "Jugador", "Amarilla", "Roja", "Total Multa"])
        
        ws_jugadores = wb.create_sheet(title="Jugadores")
        ws_jugadores.append(["Jugador", "Equipo Oficial"])
        
        wb.save(EXCEL_FILE)
    else:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if "Pagos Torneo" in wb.sheetnames:
            ws_pagos = wb["Pagos Torneo"]
            encabezados_actuales = [cell.value for cell in ws_pagos[1]]
            if encabezados_actuales != columnas_pagos:
                for col_num, header_text in enumerate(columnas_pagos, 1):
                    ws_pagos.cell(row=1, column=col_num, value=header_text)
        else:
            ws_pagos = wb.create_sheet(title="Pagos Torneo")
            ws_pagos.append(columnas_pagos)

        if "Tarjetas" not in wb.sheetnames:
            ws_tarjetas = wb.create_sheet(title="Tarjetas")
            ws_tarjetas.append(["Fecha", "Equipo", "Jugador", "Amarilla", "Roja", "Total Multa"])
            
        if "Jugadores" not in wb.sheetnames:
            ws_jugadores = wb.create_sheet(title="Jugadores")
            ws_jugadores.append(["Jugador", "Equipo Oficial"])
            
        wb.save(EXCEL_FILE)

# ----------------------------------------------------
# PROCESAMIENTO GENERAL DE MENSAJES
# ----------------------------------------------------
async def procesar_mensaje(update: Update, context: ContextTypes.DEFAULT_TYPE):
    texto_original = update.message.text.strip()
    logging.info(f"Mensaje recibido: '{texto_original}'")

    if texto_original.startswith('/'):
        return

    # MÓDULO 1: TARJETAS Y MULTAS
    if re.search(r'\b(amarilla|roja)\b', texto_original, flags=re.IGNORECASE):
        try:
            texto = texto_original
            match_fecha = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', texto)
            if match_fecha:
                d, m, y = match_fecha.groups()
                if len(y) == 2:
                    y = '20' + y
                fecha_tarjeta = f"{y}-{m.zfill(2)}-{d.zfill(2)}"
                texto = texto[match_fecha.end():].lstrip(" :-/")
            else:
                fecha_tarjeta = datetime.now().strftime("%Y-%m-%d")

            es_amarilla = bool(re.search(r'\bamarilla\b', texto, flags=re.IGNORECASE))
            tipo_texto = "Amarilla" if es_amarilla else "Roja"
            
            limpio = re.sub(r'tarjeta\s*', '', texto, flags=re.IGNORECASE)
            limpio = re.sub(r'\b(amarilla|roja)\b', '', limpio, flags=re.IGNORECASE)
            limpio = re.sub(r'\b(para|a)\b', '', limpio, count=1, flags=re.IGNORECASE)
            limpio = limpio.strip()

            partes_tarjeta = re.split(r'\s+(?:de|del|-)\s+', limpio, flags=re.IGNORECASE)
            if len(partes_tarjeta) < 2:
                raise ValueError("Falta indicar el equipo usando 'de' o 'del'")

            jugador_ingresado = partes_tarjeta[0].strip().title()
            equipo_ingresado = partes_tarjeta[1].strip().capitalize()

            inicializar_excel()
            wb = openpyxl.load_workbook(EXCEL_FILE)
            
            ws_jugadores = wb["Jugadores"]
            jugador_encontrado = False
            for row in range(2, ws_jugadores.max_row + 1):
                nombre_db = ws_jugadores.cell(row=row, column=1).value
                if nombre_db and nombre_db.lower() == jugador_ingresado.lower():
                    ws_jugadores.cell(row=row, column=2, value=equipo_ingresado)
                    jugador_encontrado = True
                    break

            if not jugador_encontrado:
                ws_jugadores.append([jugador_ingresado, equipo_ingresado])

            cant_amarilla = 1 if es_amarilla else 0
            cant_roja = 0 if es_amarilla else 1
            total_multa = (cant_amarilla * 5000) + (cant_roja * 10000)

            ws_tarjetas = wb["Tarjetas"]
            ws_tarjetas.append([fecha_tarjeta, equipo_ingresado, jugador_ingresado, cant_amarilla, cant_roja, total_multa])
            wb.save(EXCEL_FILE)
            respaldar_en_drive()

            icono = "🟨" if es_amarilla else "🟥"
            await update.message.reply_text(
                f"{icono} **Tarjeta {tipo_texto} Registrada ({fecha_tarjeta}):**\n"
                f"🛡️ Equipo: {equipo_ingresado}\n"
                f"👤 Jugador: {jugador_ingresado}\n"
                f"💵 Multa generada: ${total_multa:,}\n"
                f"☁️ *(Respaldo en Google Drive actualizado)*"
            )
            return
        except Exception as e:
            logging.error(f"Error procesando tarjeta: {e}")
            await update.message.reply_text("⚠️ Formato de tarjeta no reconocido.")
            return

    # MÓDULO 2: FINANZAS Y PAGOS
    try:
        texto = texto_original
        match_fecha = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', texto)
        if match_fecha:
            d, m, y = match_fecha.groups()
            if len(y) == 2:
                y = '20' + y
            fecha_partido = f"{y}-{m.zfill(2)}-{d.zfill(2)}"
            texto = texto[match_fecha.end():].lstrip(" :-/")
        else:
            fecha_partido = datetime.now().strftime("%Y-%m-%d")

        texto = re.sub(r'\s+', ' ', texto).strip()
        partes = re.split(r'\s+(?:contra|vs|-|\|)\s+', texto, flags=re.IGNORECASE)
        if len(partes) < 2:
            raise ValueError("Falta el conector central")

        bloque_local = partes[0].strip()
        bloque_visita = partes[1].strip()

        def extraer_pago_mixto(bloque):
            matches = list(re.finditer(r'\b(0|\d+)(?:\s*k)?\b', bloque, flags=re.IGNORECASE))
            if not matches:
                raise ValueError(f"No se encontró monto en: '{bloque}'")

            efectivo_parcial = 0
            nequi_parcial = 0
            pago_total_equipo = 0

            for i, m in enumerate(matches):
                val_str = m.group(0).lower()
                val = int(re.search(r'\d+', val_str).group())
                if 'k' in val_str and val < 1000:
                    val *= 1000
                elif 0 < val < 1000:
                    val *= 1000

                if i + 1 < len(matches):
                    siguiente_pos = matches[i+1].start()
                    segmento = bloque[m.end():siguiente_pos].lower()
                else:
                    segmento = bloque[m.end():].lower()

                if re.search(r'\b(nequi|qr|transferencia|digital)\b', segmento):
                    nequi_parcial += val
                else:
                    efectivo_parcial += val

                pago_total_equipo += val

            if efectivo_parcial == 0 and nequi_parcial == 0:
                efectivo_parcial = pago_total_equipo

            texto_equipo = bloque
            for m in matches:
                texto_equipo = texto_equipo.replace(m.group(0), "")
            
            texto_equipo = re.sub(r'\b(pag[oó]|pago|el|que|efectivo|nequi|qr|pesos|\$|y)\b', '', texto_equipo, flags=re.IGNORECASE)
            equipo = re.sub(r'\s+', ' ', texto_equipo).strip().capitalize()
            return equipo, efectivo_parcial, nequi_parcial, pago_total_equipo

        equipo_local, efec_local, neq_local, total_local = extraer_pago_mixto(bloque_local)
        equipo_visitante, efec_visita, neq_visita, total_visitante = extraer_pago_mixto(bloque_visita)

        efectivo_total_partido = efec_local + efec_visita
        nequi_total_partido = neq_local + neq_visita
        ingreso_total = total_local + total_visitante
        descuento_fijo = 70000
        caja_neta = ingreso_total - descuento_fijo

    except Exception as e:
        logging.error(f"Error procesando partido: {e}")
        await update.message.reply_text(f"⚠️ Formato financiero no reconocido. Detalle: {e}")
        return

    inicializar_excel()
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["Pagos Torneo"]
        nueva_fila = [
            fecha_partido, equipo_local, total_local, equipo_visitante, total_visitante,
            efectivo_total_partido, nequi_total_partido, ingreso_total, descuento_fijo, caja_neta
        ]
        ws.append(nueva_fila)
        wb.save(EXCEL_FILE)
        respaldar_en_drive()
    except PermissionError:
        await update.message.reply_text("⚠️ No se pudo guardar: Tienes el archivo Excel abierto en tu PC.")
        return

    await update.message.reply_text(
        f"✅ **Finanzas Registradas ({fecha_partido}):**\n"
        f"🏠 {equipo_local}: ${total_local:,} (Efectivo: ${efec_local:,} | Nequi: ${neq_local:,})\n"
        f"✈️ {equipo_visitante}: ${total_visitante:,} (Efectivo: ${efec_visita:,} | Nequi: ${neq_visita:,})\n"
        f"💵 **Efectivo Total**: ${efectivo_total_partido:,} | 📱 **Nequi Total**: ${nequi_total_partido:,}\n"
        f"💰 **Ingreso Total**: ${ingreso_total:,} | **Caja Neta**: ${caja_neta:,}\n"
        f"☁️ *(Respaldo en Google Drive actualizado)*"
    )

# ----------------------------------------------------
# FUNCIÓN DE CONFIGURACIÓN COMPARTIDA
# ----------------------------------------------------
def get_application():
    TOKEN = "8985056315:AAH0o54VgAyuMcf1qoXUXf_Okj7q2L4DvKk"
    app = ApplicationBuilder().token(TOKEN).build()
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, procesar_mensaje))
    return app

def main():
    app = get_application()
    print("🤖 Bot unificado de Torneo (Tarjetas + Pagos Mixtos) listo y escuchando en LOCAL...")
    app.run_polling()

if __name__ == '__main__':
    main()